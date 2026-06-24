import pandas as pd
import io
import re
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment

# ============================================================
# CONFIGURACIÓN
# ============================================================
INTERVALO_MIN      = 5
UMBRAL_IRRADIANCIA = 1.0
UMBRAL_PERDIDA     = 1.0

# Función auxiliar para extraer fecha del nombre del archivo
def extraer_fecha(nombre_archivo: str) -> str:
    """Fallback: extrae fecha del nombre del archivo."""
    patrones = [
        (r'(\d{4})[_\-](\d{2})[_\-](\d{2})', '%Y-%m-%d'),
        (r'(\d{4})(\d{2})(\d{2})',             '%Y%m%d'),
        (r'(\d{2})[_\-](\d{2})[_\-](\d{4})',  '%d-%m-%Y'),
    ]
    for patron, fmt in patrones:
        m = re.search(patron, nombre_archivo)
        if m:
            candidato = '-'.join(m.groups())
            try:
                return datetime.strptime(candidato, fmt).strftime('%d/%m/%Y')
            except ValueError:
                continue
    return 'Fecha desconocida'

# Función principal para extraer la fecha
def extraer_fecha_csv(contenido: str, nombre_archivo: str) -> str:
    """
    Extrae la fecha de la primera línea del CSV.
    Formato esperado: 'Period: 9/5/2026 0:00:00 - ...'
    Si no la encuentra, hace fallback al nombre del archivo.
    """
    primera_linea = contenido.splitlines()[0] if contenido else ''
    m = re.search(r'Period:\s*(\d{1,2})/(\d{1,2})/(\d{4})', primera_linea)
    if m:
        dia, mes, anio = m.group(1), m.group(2), m.group(3)
        try:
            return datetime(int(anio), int(mes), int(dia)).strftime('%d/%m/%Y')
        except ValueError:
            pass
    return extraer_fecha(nombre_archivo)


def cargar_datos(archivo) -> tuple[pd.DataFrame, str]:
    """
    Acepta un InMemoryUploadedFile de Django.
    Retorna (DataFrame recortado a ventana operativa, fecha extraída del CSV).
    """
    try:
        contenido = archivo.read().decode('utf-16', errors='ignore').replace('"', '')
    except Exception as e:
        raise RuntimeError(f"No se pudo leer el archivo: {e}")

    fecha = extraer_fecha_csv(contenido, archivo.name)

    df = pd.read_csv(
        io.StringIO(contenido),
        sep='\t',
        skiprows=3,
        header=None,
        engine='python'
    )

    if df.shape[1] < 3:
        raise ValueError(
            f"Formato inesperado: se esperaban >= 3 columnas, "
            f"se encontraron {df.shape[1]}."
        )

    df = df.iloc[:, :3].copy()
    df.columns = ['Hora', 'Power', 'Irradiancia']
    df['Power']       = pd.to_numeric(df['Power'],       errors='coerce')
    df['Irradiancia'] = pd.to_numeric(df['Irradiancia'], errors='coerce')
    df = df.dropna(subset=['Power', 'Irradiancia'], how='all').reset_index(drop=True)

    potencia_max = df['Power'].max(skipna=True)
    if potencia_max > 1000:
        df['Power'] = df['Power'] / 1000

    operativo = df['Irradiancia'] >= UMBRAL_IRRADIANCIA
    if not operativo.any():
        return pd.DataFrame(), fecha

    primer_idx = operativo.idxmax()
    ultimo_idx  = operativo[::-1].idxmax()
    return df.loc[primer_idx:ultimo_idx].copy().reset_index(drop=True), fecha


def calcular_perdidas(df: pd.DataFrame) -> tuple[pd.DataFrame, float, float]:
    """
    Retorna (df_calculado, promedio_limpio, umbral_dinamico).
    El umbral se calcula como el 10% de la generación esperada
    por intervalo promedio del día.
    """
    factor_t = INTERVALO_MIN / 60
    df = df.copy()

    mask_irr = df['Irradiancia'] > 0
    df['Eficiencia'] = float('nan')
    df.loc[mask_irr, 'Eficiencia'] = (
        df.loc[mask_irr, 'Power'] / df.loc[mask_irr, 'Irradiancia']
    ).round(5)

    promedio_inicial = df['Eficiencia'].mean(skipna=True)

    # Umbral dinámico: 10% de la generación esperada por intervalo típico
    irradiancia_promedio = df.loc[mask_irr, 'Irradiancia'].mean()
    umbral_dinamico = round(
        promedio_inicial * irradiancia_promedio * factor_t * 0.10, 4
    )
    # Seguridad: nunca menor a 0.01 kWh para evitar falsos positivos extremos
    umbral_dinamico = max(umbral_dinamico, 0.01)

    perdida_potencial = ((df['Irradiancia'] * promedio_inicial) - df['Power']) * factor_t
    df['Error'] = (perdida_potencial >= umbral_dinamico) | (
        (df['Power'] == 0) & (df['Irradiancia'] > 0)
    )

    # Pasada 1: descartar lapsos insignificantes
    promedio_prov = df.loc[~df['Error'], 'Eficiencia'].mean(skipna=True)
    df['_prov'] = 0.0
    mask_e = df['Error']
    df.loc[mask_e, '_prov'] = (
        (df.loc[mask_e, 'Irradiancia'] * promedio_prov) - df.loc[mask_e, 'Power']
    ) * factor_t
    df['_prov'] = df['_prov'].clip(lower=0)

    errores_prov          = df[df['Error']].reset_index(drop=False).rename(columns={'index': 'idx_original'})
    errores_prov['grupo'] = (errores_prov['idx_original'].diff().fillna(1) != 1).cumsum()
    totales_lapso         = errores_prov.groupby('grupo')['_prov'].sum()
    grupos_insig          = set(totales_lapso[totales_lapso < umbral_dinamico].index)
    indices_desmarcar     = errores_prov[errores_prov['grupo'].isin(grupos_insig)]['idx_original'].values
    df.loc[indices_desmarcar, 'Error'] = False
    df = df.drop(columns=['_prov'])

    # Pasada 2: promedio y pérdidas finales
    promedio_limpio = df.loc[~df['Error'], 'Eficiencia'].mean(skipna=True)
    df['Perdida_kWh'] = 0.0
    mask_error = df['Error']
    df.loc[mask_error, 'Perdida_kWh'] = (
        (df.loc[mask_error, 'Irradiancia'] * promedio_limpio) - df.loc[mask_error, 'Power']
    ) * factor_t
    df['Perdida_kWh'] = df['Perdida_kWh'].clip(lower=0).round(4)

    return df, round(float(promedio_limpio), 6), umbral_dinamico


def agrupar_lapsos(df: pd.DataFrame, umbral_dinamico: float) -> pd.DataFrame:
    errores = df[df['Error']].copy()
    if errores.empty:
        return pd.DataFrame(
            columns=['Hora_Inicio', 'Hora_Fin', 'Intervalos', 'Perdida_Total_kWh']
        )

    errores               = errores.reset_index(drop=False).rename(columns={'index': 'idx_original'})
    errores['grupo']      = (errores['idx_original'].diff().fillna(1) != 1).cumsum()

    lapsos = (
        errores.groupby('grupo')
        .agg(
            Hora_Inicio       = ('Hora',        'first'),
            Hora_Fin          = ('Hora',        'last'),
            Intervalos        = ('Hora',        'count'),
            Perdida_Total_kWh = ('Perdida_kWh', 'sum'),
        )
        .reset_index(drop=True)
    )
    lapsos['Perdida_Total_kWh'] = lapsos['Perdida_Total_kWh'].round(4)
    return lapsos[lapsos['Perdida_Total_kWh'] >= umbral_dinamico].reset_index(drop=True)


def generar_excel(registros: list[dict]) -> bytes:
    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = 'Resumen de Pérdidas'

    fill_header = PatternFill('solid', fgColor='2B6CB0')
    fill_fila   = PatternFill('solid', fgColor='FEE2E2')
    fill_total  = PatternFill('solid', fgColor='E2E8F0')
    font_header = Font(bold=True, color='FFFFFF', size=11)
    font_normal = Font(size=10)
    font_total  = Font(bold=True, size=11)
    centro      = Alignment(horizontal='center', vertical='center')

    def c(fila, col, valor, font=None, fill=None):
        cell           = ws.cell(row=fila, column=col, value=valor)
        cell.font      = font or font_normal
        cell.fill      = fill or PatternFill()
        cell.alignment = centro
        return cell

    headers = ['Fecha', 'Hora Inicio', 'Hora Fin', 'Pérdida (kWh)']
    for col, h in enumerate(headers, start=1):
        c(1, col, h, font=font_header, fill=fill_header)

    for fila, reg in enumerate(registros, start=2):
        c(fila, 1, reg['fecha'])
        c(fila, 2, reg['hora_inicio'])
        c(fila, 3, reg['hora_fin'])
        c(fila, 4, reg['perdida_total_kWh'], fill=fill_fila)

    fila_total = len(registros) + 2
    ws.merge_cells(f'A{fila_total}:D{fila_total}')
    c(fila_total, 1, 'PÉRDIDA TOTAL', font=font_total, fill=fill_total)
    total = round(sum(r['perdida_total_kWh'] for r in registros), 2)
    c(fila_total, 5, total, font=font_total, fill=fill_total)

    for col, ancho in zip('ABCDE', [14, 28, 14, 14, 16]):
        ws.column_dimensions[col].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def procesar_archivos(archivos) -> dict:
    resultados_vista   = []
    registros_excel    = []
    archivos_omitidos  = []
    archivos_con_error = []

    for archivo in archivos:
        nombre = archivo.name

        try:
            df_crudo, fecha = cargar_datos(archivo)

            if df_crudo.empty:
                archivos_omitidos.append(nombre)
                continue

            df_calc, promedio_limpio, umbral_dinamico = calcular_perdidas(df_crudo)
            lapsos                   = agrupar_lapsos(df_calc, umbral_dinamico)
            perdida_total_dia        = round(float(lapsos['Perdida_Total_kWh'].sum()), 2)

            resultados_vista.append({
                'nombre':        nombre,
                'fecha':         fecha,
                'umbral':        umbral_dinamico,
                'promedio':      promedio_limpio,
                'perdida_total': perdida_total_dia,
                'sin_perdidas':  lapsos.empty,
                'lapsos':        lapsos.to_dict('records') if not lapsos.empty else [],
            })

            for _, lapso in lapsos.iterrows():
                registros_excel.append({
                    'fecha':             fecha,
                    'nombre_archivo':    nombre,
                    'hora_inicio':       lapso['Hora_Inicio'],
                    'hora_fin':          lapso['Hora_Fin'],
                    'perdida_total_kWh': lapso['Perdida_Total_kWh'],
                })

        except Exception as e:
            archivos_con_error.append({'nombre': nombre, 'error': str(e)})
            continue

    excel_bytes    = generar_excel(registros_excel) if registros_excel else None
    perdida_global = round(sum(r['perdida_total'] for r in resultados_vista), 2)

    return {
        'resultados':          resultados_vista,
        'registros_excel':     registros_excel,
        'excel_bytes':         excel_bytes,
        'archivos_omitidos':   archivos_omitidos,
        'archivos_con_error':  archivos_con_error,
        'perdida_global':      perdida_global,
        'hay_resultados':      bool(resultados_vista),
    }